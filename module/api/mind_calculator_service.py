"""心智单元计算器的实例数据与文件导入导出接口。"""
import base64
import binascii
import csv
import io
import json
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError

from pydantic import ValidationError

from deploy.atomic import atomic_write
from module.api.protocol import ApiError, MindShip
from module.config.transaction import config_transaction
from module.runtime.mind_calculator import RARITY_NAMES, calculate, catalog, revision

TASK, GROUP = 'MindCalculatorScan', 'MindCalculator'


def input_ship(ship):
    """规范化用户输入的舰船字段并校验类型。"""
    return MindShip.model_validate(ship).model_dump()


def decode(content):
    """解码导入文件的内容并检查数据格式。"""
    try:
        return base64.b64decode(content, validate=True)
    except (ValueError, binascii.Error):
        raise ApiError('INVALID_PARAMS', '文件内容不是有效的 Base64') from None


def table_ships(rows):
    """兼容原计算器的舰船明细表头，不依据位置猜测列含义。"""
    aliases = {'name': ('name', '船名', '舰船', '名称'), 'level': ('level', '等级'),
               'rarity': ('rarity', '稀有度'), 'base_rarity': ('base_rarity', '基础稀有度'),
               'excluded': ('excluded', '排除'), 'review': ('review', '需核对'), 'source': ('source', '来源'),
               'status': ('status', '状态'), 'notes': ('notes', '备注')}
    columns = None
    output = []
    for index, row in enumerate(rows):
        if index > 5020:
            raise ValueError('最多导入 5000 艘舰船')
        if columns is None:
            header = [str(cell or '').strip() for cell in row]
            columns = {key: next((header.index(name) for name in names if name in header), -1)
                       for key, names in aliases.items()}
            if columns['name'] < 0 or columns['level'] < 0:
                columns = None
                if index >= 19:
                    break
                continue
            continue
        def value(key, default=''):
            """读取导入单元格并转为可用的字符串。"""
            column = columns[key]
            return row[column] if 0 <= column < len(row) and row[column] is not None else default
        name = str(value('name')).strip()
        if len(name) > 1 and name.startswith("'") and name[1:2] in '=+-@\t\r':
            name = name[1:]
        if not name:
            continue
        raw_level = value('level', 0)
        if raw_level == '':
            raw_level = 0
        if isinstance(raw_level, bool) or str(raw_level).strip() not in {str(x) for x in range(126)}:
            raise ValueError(f'第 {index + 1} 行等级必须为 0 至 125 的整数')
        rarity = str(value('rarity')).strip()
        rarity = {cn: key for key, cn in RARITY_NAMES.items()}.get(rarity, rarity.upper())
        base_rarity = str(value('base_rarity')).strip()
        base_rarity = {cn: key for key, cn in RARITY_NAMES.items()}.get(base_rarity, base_rarity.upper())
        def boolean(key):
            """将导入的标记字段转换为布尔值。"""
            raw = str(value(key)).strip().casefold()
            if raw not in ('', '0', '1', 'true', 'false', '是', '否'):
                raise ValueError(f'第 {index + 1} 行 {key} 字段无效')
            return raw in ('1', 'true', '是')
        output.append(input_ship(dict(name=name, level=int(raw_level), rarity=rarity,
                                      base_rarity=base_rarity, excluded=boolean('excluded'),
                                      review=boolean('review') or value('status') in ('存疑', '等级未读到', '名字未匹配'),
                                      source=str(value('source') or value('notes'))[:200])))
    if columns is None:
        raise ValueError('缺少“船名”和“等级”表头')
    return output


def import_ships(filename, content):
    """从上传的文件中解析舰船列表。"""
    raw = decode(content)
    suffix = Path(filename).suffix.casefold()
    try:
        if suffix == '.json':
            data = json.loads(raw.decode('utf-8-sig'))
            items = data.get('ships') if isinstance(data, dict) else data
            if not isinstance(items, list) or len(items) > 5000:
                raise ValueError('JSON 必须含 ships 数组，最多 5000 艘')
            return [input_ship(ship) for ship in items]
        if suffix == '.csv':
            return table_ships(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError('Excel 解压后超过 32 MiB 限制')
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            try:
                sheet = book['舰船明细'] if '舰船明细' in book.sheetnames else book.active
                rows = list(sheet.iter_rows(max_row=5021, max_col=20, values_only=True))
                ships = table_ships(rows)
                # 原工具把排除清单放在明细右侧 M 列；按明确表头回读。
                for index, row in enumerate(rows[:20]):
                    if '排除船名' in row:
                        column = row.index('排除船名')
                        for values in rows[index + 1:]:
                            name = values[column]
                            if name and not str(name).startswith('本次没有需要排除的船'):
                                ships.append(input_ship(dict(name=str(name), level=0, excluded=True)))
                        break
                if len(ships) > 5000:
                    raise ValueError('最多导入 5000 艘舰船')
                return ships
            finally:
                book.close()
        raise ValueError('请选择 JSON、CSV 或 XLSX 文件')
    except (ValueError, TypeError, KeyError, ValidationError, zipfile.BadZipFile, OSError, ParseError) as exc:
        raise ApiError('INVALID_PARAMS', f'导入失败：{exc}') from None


class MindCalculatorService:
    def __init__(self, configs):
        """绑定配置服务并初始化心智计算器接口。"""
        self.configs = configs

    def report(self, instance):
        """返回指定实例保存的舰船列表及版本信息。"""
        data, _ = self.configs.read(instance)
        saved = data.get(TASK, {}).get(GROUP, {}).get('Result', {})
        ships = saved.get('ships', [])
        fields = data.get(TASK, {}).get(GROUP, {})
        return dict(instance=instance, revision=revision(ships), updated_at=saved.get('updated_at', ''),
                    min_level=fields.get('MinLevel', 95), max_level=fields.get('MaxLevel', 120),
                    **calculate(ships))

    def save(self, instance, expected_revision, ships):
        """按版本校验保存实例的心智计算器清单。"""
        with self.configs.lock, config_transaction(self.configs.path(instance)):
            data, _ = self.configs.read(instance)
            fields = data.setdefault(TASK, {}).setdefault(GROUP, {})
            if revision(fields.get('Result', {}).get('ships', [])) != expected_revision:
                raise ApiError('CONFLICT', '舰船数据已变化，请重新载入后再保存')
            fields['Result'] = dict(ships=ships, updated_at=datetime.now().isoformat(timespec='seconds'))
            atomic_write(str(self.configs.path(instance)), json.dumps(data, ensure_ascii=False, indent=2))
            return self.report(instance)

    def catalog(self, instance):
        """返回用于计算与人工核对的舰船资料。"""
        self.configs.path(instance)
        data = catalog()
        return dict(updated_at=data['updated_at'], ships=list(data['ships'].values()))

    def calculate(self, instance, ships):
        """计算选定舰船升级所需的心智单元。"""
        self.configs.path(instance)
        return calculate(ships)

    def import_file(self, instance, filename, content):
        """解析用户上传文件并返回可检查的舰船记录。"""
        self.configs.path(instance)
        return {'ships': import_ships(filename, content)}

    def recognize(self, instance, filename, content):
        """识别上传截图中的舰船卡片和等级。"""
        self.configs.path(instance)
        from PIL import Image, UnidentifiedImageError
        from module.runtime.mind_calculator import recognize
        try:
            raw = decode(content)
            if len(raw) > 5_500_000:
                raise ValueError('图片超过 5.5 MB 限制，请压缩后重试')
            with Image.open(io.BytesIO(raw)) as image:
                if image.format not in ('PNG', 'JPEG', 'WEBP'):
                    raise ValueError('图片格式不支持，请上传 PNG、JPEG 或 WebP')
                ships = recognize(image, Path(filename.replace('\\', '/')).name)
        except (ValueError, OSError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
            raise ApiError('INVALID_PARAMS', f'截图识别失败：{exc}') from None
        except RuntimeError as exc:
            raise ApiError('OCR_UNAVAILABLE', f'OCR 引擎无法完成识别，请检查模型和 OCR 配置：{str(exc)[:200]}') from None
        return {'ships': ships}

    def export(self, instance, format):
        """将指定实例的舰船清单导出为选定格式。"""
        report = self.report(instance)
        headers = ['船名', '等级', '稀有度', '基础稀有度', '排除', '需核对', '来源']
        fields = ['name', 'level', 'rarity', 'base_rarity', 'excluded', 'review', 'source']
        ships = [{key: ship[key] for key in fields} for ship in report['ships']]
        if format == 'json':
            raw = json.dumps(dict(version=1, ships=ships), ensure_ascii=False, indent=2).encode()
        elif format == 'csv':
            output = io.StringIO(newline='')
            writer = csv.writer(output)
            writer.writerow(headers)
            def csv_value(value):
                """安全编码 CSV 单元格以防公式注入。"""
                text = str(value)
                return "'" + text if text[:1] in ('=', '+', '-', '@', '\t', '\r') else text
            writer.writerows([[csv_value(ship[key]) for key in fields] for ship in ships])
            raw = output.getvalue().encode('utf-8-sig')
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            book = Workbook()
            summary = book.active
            summary.title = '心智单元汇总'
            summary.append(['稀有度', '≤100', '101–105', '106–110', '111–115', '116–119', '≥120', '合并后舰船', '心智单元Ⅰ', '物资'])
            for bucket in report['summary']:
                summary.append([RARITY_NAMES[bucket['rarity']], *bucket['stages'], bucket['count'], bucket['mind'], bucket['gold']])
            summary.append(['合计', *[sum(bucket['stages'][i] for bucket in report['summary']) for i in range(6)],
                            report['included'], report['mind'], report['gold']])
            summary.append(['目标 120 级；不含 120→125；改造按基础稀有度；同名只计最高等级；待核对数据不计费。'])
            for title, subset in [('舰船明细', ships), ('需核对清单', [ship for ship in ships if ship['review'] or not ship['level']])]:
                sheet = book.create_sheet(title)
                sheet.append(headers)
                for ship in subset:
                    sheet.append([ship[key] for key in fields])
            for sheet in book:
                sheet.freeze_panes = 'B2'
                sheet.auto_filter.ref = sheet.dimensions
                for column in sheet.columns:
                    sheet.column_dimensions[column[0].column_letter].width = 20
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='2F80ED')
                # 名称和来源可能来自上传文件，显式保存成文本，避免公式执行。
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type == 'f':
                            cell.data_type = 's'
            output = io.BytesIO()
            book.save(output)
            raw = output.getvalue()
        return dict(filename=f'{instance}-mind-calculator.{format}', content=base64.b64encode(raw).decode())
